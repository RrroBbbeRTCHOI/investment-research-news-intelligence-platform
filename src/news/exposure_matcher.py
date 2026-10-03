"""Exact channel + country joins with evidence and point-in-time eligibility."""
from datetime import date, datetime, time, timezone
from pathlib import Path
import math
import re
from .constants import (CHANNEL_FAMILIES, COUNTRY_ALIASES, EVIDENCE_QUALITY, MAG7, WEIGHTS,
                        ENTITY_ALIASES, SERVICE_ALIASES, SERVICE_OWNERS, DEPENDENCY_ROLES,
                        PRODUCTION_EDGE_CHANNELS)

WORKBOOK_NAME = 'MAG7_News_Intelligence_Exposure_Knowledge_Layer_V1.xlsx'
REQUIRED_COLUMNS = ('edge_id','ticker','country','channel_family','channel_subtype',
 'role','product_or_business','counterparty','facility_or_route','commodity_or_input',
 'stage','exposure_value','exposure_unit','exposure_band','evidence_status','source_url',
 'source_locator','valid_from','valid_to','known_at_utc','confidence','notes')

def parse_time(value, *, end_of_day=False) -> datetime | None:
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        dt = value
    elif isinstance(value, date):
        dt = datetime.combine(value, time.max if end_of_day else time.min)
    else:
        text = str(value).strip()
        if len(text) == 10:
            dt = datetime.combine(date.fromisoformat(text), time.max if end_of_day else time.min)
        else:
            dt = datetime.fromisoformat(text.replace('Z','+00:00'))
    # Spreadsheet date/time cells and ISO values without offsets use documented UTC.
    return dt.replace(tzinfo=timezone.utc) if dt.tzinfo is None else dt.astimezone(timezone.utc)

def load_exposure_edges(root: Path, workbook: Path | None = None) -> tuple[list[dict], dict]:
    if workbook is None:
        found = sorted(p for p in root.rglob(WORKBOOK_NAME) if not any(x.startswith('.') for x in p.relative_to(root).parts))
        if len(found) > 1:
            return [], dict(status='ambiguous_workbook', message='Multiple exposure workbooks found; use --exposure-workbook.', edge_count=0)
        workbook = found[0] if found else None
    if workbook is None or not workbook.is_file():
        return [], dict(status='missing_workbook', message=f'{WORKBOOK_NAME} is missing; zero knowledge-layer matches.', edge_count=0)
    try:
        from openpyxl import load_workbook
    except ImportError:
        return [], dict(status='missing_reader', message='Install src/news/requirements.txt to read the supplied exposure workbook.', edge_count=0)
    from zipfile import BadZipFile
    from xml.etree.ElementTree import ParseError
    from openpyxl.utils.exceptions import InvalidFileException
    try:
        wb = load_workbook(workbook, read_only=True, data_only=True)
        try:
            if 'Exposure Edges' not in wb.sheetnames:
                return [], dict(status='invalid_workbook', message='Missing Exposure Edges sheet.', edge_count=0)
            rows = wb['Exposure Edges'].iter_rows(values_only=True)
            # Permit a title before the actual table; require all frozen columns.
            header = None
            for row in rows:
                values = [str(v).strip() if v is not None else '' for v in row]
                if 'edge_id' in values and 'ticker' in values:
                    header = values
                    break
            if not header or not set(REQUIRED_COLUMNS) <= set(header):
                return [], dict(status='invalid_workbook', message='Exposure Edges is missing required columns.', edge_count=0)
            if len([x for x in header if x]) != len(set(x for x in header if x)):
                return [], dict(status='invalid_workbook', message='Duplicate Exposure Edges columns.', edge_count=0)
            edges = []
            for row in rows:
                if not any(v is not None for v in row):
                    continue
                edge = dict(zip(header,row))
                edge = {k: (v.isoformat() if isinstance(v,(date,datetime)) else v.strip() if isinstance(v,str) else v)
                        for k,v in edge.items() if k in REQUIRED_COLUMNS}
                edges.append(edge)
        finally:
            wb.close()
    except (OSError, BadZipFile, ParseError, InvalidFileException, ValueError) as exc:
        return [], dict(status='invalid_workbook', message=f'Cannot read exposure workbook ({type(exc).__name__}).', edge_count=0)
    ids = [e.get('edge_id') for e in edges]
    if len(ids) != len(set(ids)):
        return [], dict(status='invalid_workbook', message='Duplicate edge IDs; no edges loaded.', edge_count=0)
    return edges, dict(status='loaded', workbook=workbook.name, edge_count=len(edges),
                        message='Loaded raw exposure rows; each match still requires evidence and temporal checks.')

def edge_ineligibility(edge: dict, prediction_time) -> str | None:
    if not edge.get('edge_id') or edge.get('ticker') not in MAG7 or edge.get('channel_family') not in CHANNEL_FAMILIES:
        return 'invalid_edge_identity_or_channel'
    if edge.get('evidence_status') not in EVIDENCE_QUALITY:
        return 'blocked_or_hypothesis_evidence'
    if not edge.get('source_url') or not edge.get('source_locator'):
        return 'missing_source_provenance'
    stage = str(edge.get('stage') or '').lower()
    role = str(edge.get('role') or '').lower()
    if any(token in stage+' '+role for token in ('planned','proposed','construction','announced','pre-operat','pilot','future')):
        return 'non_operating_capacity'
    # Facility capacity needs an affirmative operating stage, not merely absence of the word planned.
    if edge.get('facility_or_route') and stage not in ('operating','operational','commercial','commercial_operation','in_production'):
        return 'unverified_operating_stage'
    try:
        prediction = parse_time(prediction_time)
        known = parse_time(edge.get('known_at_utc'),end_of_day=True)
        start = parse_time(edge.get('valid_from'))
        end = parse_time(edge.get('valid_to'),end_of_day=True)
    except (TypeError,ValueError,OverflowError):
        return 'invalid_timestamp'
    if prediction is None:
        return 'missing_prediction_time'
    if known is None:
        return 'missing_known_at_unsuitable_for_history'
    if start is None:
        return 'missing_valid_from'
    if end is not None and end < start:
        return 'invalid_validity_interval'
    if known > prediction:
        return 'future_known_at'
    if start > prediction:
        return 'future_valid_from'
    if end is not None and prediction > end:
        return 'expired_edge'
    return None

def _contains(text: str, term: str) -> bool:
    return bool(re.search(r'(?<!\w)'+re.escape(term)+r'(?!\w)',text,re.I))

def _event_targets(event: dict) -> tuple[set[str], dict[str, tuple[str, ...]], str]:
    """Operational subject identities; no new frozen event fields or inferred edges."""
    aliases = dict(ENTITY_ALIASES)
    known_aliases = {a.lower() for values in aliases.values() for a in values}
    for name in event.get('companies_mentioned',[]):
        if name.lower() not in known_aliases:
            aliases[name] = (name,)
    for text in (event.get('event_summary') or '',event.get('canonical_headline') or ''):
        mentions=[]
        for entity, names in aliases.items():
            for name in names:
                for m in re.finditer(r'(?<!\w)'+re.escape(name)+r'(?!\w)',text,re.I):
                    mentions.append((m.start(),m.end(),entity))
        if not mentions:
            continue
        trigger=re.search(r'\b(?:outage|shutdown|production|factory|cyberattack|ransomware|breach|earnings|results|financing|funding|CEO|chief executive|failure)\b',text,re.I)
        pivot=trigger.start() if trigger else len(text)
        before=[m for m in mentions if m[1]<=pivot]
        chosen=max(before,key=lambda m:m[1]) if before else min(mentions,key=lambda m:abs(m[0]-pivot))
        targets={chosen[2]}
        # Coordinated subjects (AWS and Azure) are allowed; a quoted observer is not.
        for start,end,entity in mentions:
            if end<=chosen[0] and re.fullmatch(r'[\s,&]*(?:and|or)?[\s,&]*',text[end:chosen[0]],re.I):
                targets.add(entity)
        return targets,aliases,text
    return set(),aliases,event.get('event_summary') or ''

def semantic_edge_reason(event: dict, edge: dict) -> str | None:
    """Return an exclusion reason before any exposure component can score."""
    operational = event.get('event_type') in {
        'infrastructure_outage','supply_chain','cybersecurity','corporate_financing',
        'corporate_earnings','management_governance','product_technology'}
    if not operational:
        return None
    targets,aliases,text = _event_targets(event)
    regional_grid = bool(re.search(r'\b(?:regional|national|statewide|citywide)\b.{0,45}\b(?:grid|electricity|power)\b',text,re.I))
    if event.get('event_subtype') in ('port_disruption','transport_disruption') or regional_grid:
        return None
    if not targets:
        # An unnamed operational subject is not permission to match every peer operator.
        if re.search(r'\bgrid\b',text,re.I) and event.get('power_disruption') is True:
            return None
        return 'operational_subject_unknown'
    role=str(edge.get('role') or '').lower()
    if event.get('event_type') in ('infrastructure_outage','supply_chain','cybersecurity') and (
            role=='seller' or edge.get('stage')=='demand'):
        return 'operational_event_not_demand_exposure'
    if edge.get('ticker') in targets:
        business=str(edge.get('product_or_business') or '')
        for service,names in SERVICE_ALIASES.items():
            if (SERVICE_OWNERS[service] == edge.get('ticker') and any(_contains(text,a) for a in names)
                    and business and not any(_contains(business,a) for a in names)):
                return 'different_product_or_business'
        return None
    if role not in DEPENDENCY_ROLES:
        return 'entity_specific_without_upstream_dependency_role'
    # A relation field must explicitly connect this company to the affected subject.
    # Broad product/industry descriptions and shared country are never such a link.
    relationships = ' ; '.join(str(edge.get(k) or '') for k in
                              ('counterparty','facility_or_route','commodity_or_input'))
    specific_services = [names for service,names in SERVICE_ALIASES.items()
                         if SERVICE_OWNERS[service] in targets and any(_contains(text,a) for a in names)]
    relation_aliases = [a for names in specific_services for a in names] if specific_services else [
        name for target in targets for name in aliases[target]]
    if any(_contains(relationships,name) for name in relation_aliases):
        return None
    location = event.get('location_name')
    if (location and event.get('evidence_fields',{}).get('location_name') and
            _contains(str(edge.get('facility_or_route') or ''),location)):
        return None
    if event.get('supply_shortage') is True:
        if any(_contains(str(edge.get('commodity_or_input') or ''),commodity)
               for commodity in event.get('commodities_mentioned',[])):
            return None
    return 'entity_specific_without_shared_dependency'

def match_exposures(event: dict, candidates: dict, edges: list[dict], *, prediction_time=None) -> tuple[list[dict], list[dict]]:
    if candidates.get('event_id') != event.get('event_id'):
        raise ValueError('Candidate channels belong to another event')
    countries = {COUNTRY_ALIASES.get(c,c) for c in event.get('affected_countries',[])}
    if event.get('primary_country'):
        countries.add(COUNTRY_ALIASES.get(event['primary_country'],event['primary_country']))
    channels = [c for c in candidates.get('candidate_channels',[]) if c.get('status') == 'candidate']
    rejected, eligible = [], []
    for edge in edges:
        reason = edge_ineligibility(edge,prediction_time)
        if reason:
            rejected.append(dict(event_id=event['event_id'], edge_id=edge.get('edge_id'), reason=reason))
        else:
            eligible.append(edge)
    pairs = []
    for ticker in MAG7:
        matched = []
        for edge in eligible:
            if edge['ticker'] != ticker:
                continue
            if COUNTRY_ALIASES.get(edge.get('country'),edge.get('country')) not in countries:
                continue
            # Unspecified candidates cannot prove a narrow subchannel. A generic edge accepts any same-family candidate.
            compatible = any(c['channel_family'] == edge['channel_family'] and
                (not edge.get('channel_subtype') or edge['channel_subtype'] in ('unspecified','general') or
                 c.get('channel_subtype') == edge['channel_subtype']) for c in channels)
            refined = False
            if not compatible and event.get('event_type')=='supply_chain' and event.get('production_disruption') is True:
                targets,aliases,_ = _event_targets(event)
                # A named supplier relation may refine a broad production hypothesis.
                # Country/industry overlap alone can never take this path.
                refined = (edge.get('role') in DEPENDENCY_ROLES and edge['ticker'] not in targets
                    # Explicit process evidence limits supplier-name refinement to
                    # that process; fabrication does not imply packaging or vice versa.
                    and not any(c['channel_family'] in ('advanced_semiconductors', 'advanced_packaging')
                                for c in channels)
                    and edge['channel_family'] in PRODUCTION_EDGE_CHANNELS
                    and edge.get('channel_subtype') in (None,'','unspecified','general')
                    and any(c['channel_family'] in ('manufacturing','supply_chain') for c in channels)
                    and any(_contains(str(edge.get('counterparty') or ''),name)
                            for target in targets for name in aliases[target]))
                compatible = refined
            if not compatible:
                continue
            semantic_reason = semantic_edge_reason(event,edge)
            if semantic_reason:
                rejected.append(dict(event_id=event['event_id'],edge_id=edge.get('edge_id'),reason=semantic_reason))
                continue
            reasons=['country_match','channel_match','eligible_evidence','point_in_time_valid','semantic_compatibility']
            if refined:
                reasons.append('named_supplier_production_channel_refinement')
            matched.append({**edge, 'match_reasons':reasons})
        direct = ticker in event.get('tickers_mentioned',[])
        quality = max((EVIDENCE_QUALITY[e['evidence_status']] for e in matched),default=0.0)
        severity = event.get('severity_score')
        if severity is not None and (type(severity) not in (int,float) or not math.isfinite(severity) or not 0<=severity<=1):
            raise ValueError('Invalid event severity')
        components = dict(direct_mention=float(direct),edge_match=float(bool(matched)),
            country_match=float(bool(matched)),channel_match=float(bool(matched)),
            evidence_quality=quality,event_severity=severity)
        # All exposure terms are gated by a qualifying edge. Unknown severity contributes no bonus.
        score = WEIGHTS['direct_mention']*direct
        if matched:
            score += quality*(WEIGHTS['edge_match']+WEIGHTS['country_match']+WEIGHTS['channel_match']+
                              WEIGHTS['evidence_quality']+WEIGHTS['event_severity']*(severity or 0))
        pairs.append(dict(event_id=event['event_id'],ticker=ticker,matched=bool(matched),
            relevance_score=round(min(1.0,max(0.0,score)),3),matched_edges=matched,direct_mention=direct,
            relevance_components=components,model_status='heuristic_v1',prediction_time=prediction_time))
    return pairs,rejected
